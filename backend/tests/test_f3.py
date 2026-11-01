"""F3: alertas ambientais, relatórios/exportações, qualidade de dados e sugestões de IA."""

import io
import json

import pytest
import responses
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import load_workbook

from itree.alertas import fontes
from itree.alertas.models import AlertaAmbiental, IngestaoAlertas
from itree.alertas.tasks import ingerir_fontes_configuradas
from itree.inventario.models import Arvore, Especie, Medida
from itree.ocorrencias.models import Ocorrencia
from itree.qualidade.models import InferenciaIA
from itree.relatorios.models import Exportacao

from .conftest import PONTO_A, PONTO_B, com_escopo, criar_usuario

pytestmark = pytest.mark.django_db

FOCOS = (
    "id,lat,lon,data_hora_gmt,satelite,municipio\n"
    "F1,-3.755,-38.555,2026-10-30 17:20:00,AQUA_M-T,FORTALEZA\n"  # Bairro A
    "F2,-3.755,-38.535,2026-10-30 17:25:00,AQUA_M-T,FORTALEZA\n"  # Bairro B
    "F3,-5.100,-39.900,2026-10-30 17:30:00,AQUA_M-T,OUTRO\n"  # fora do município
    "F4,abc,-38.5,2026-10-30 17:30:00,AQUA_M-T,FORTALEZA\n"  # inválido
)


def poligono_utm():
    """Polígono de 100 m x 100 m (1 ha) em SIRGAS 2000 / UTM 24S, dentro do Bairro A."""
    from django.contrib.gis.geos import Point

    centro = Point(-38.555, -3.755, srid=4326)
    centro.transform(31984)
    x, y = centro.x, centro.y
    anel = [[x, y], [x + 100, y], [x + 100, y + 100], [x, y + 100], [x, y]]
    return json.dumps(
        {
            "type": "FeatureCollection",
            "crs": {"type": "name", "properties": {"name": "urn:ogc:def:crs:EPSG::31984"}},
            "features": [
                {
                    "type": "Feature",
                    "properties": {
                        "gid": 77,
                        "view_date": "2026-10-20",
                        "classname": "DESMATAMENTO_CR",
                    },
                    "geometry": {"type": "Polygon", "coordinates": [anel]},
                }
            ],
        }
    ).encode()


@pytest.fixture
def analista(territorios):
    return com_escopo(
        criar_usuario("ana@itree.teste", ["analista_ambiental"]), territorios["municipio"]
    )


# --- Ingestão -----------------------------------------------------------------


def test_ingestao_de_focos_idempotente_e_restrita_ao_municipio(territorios):
    r = fontes.ingerir("inpe_focos", FOCOS.encode(), "focos.csv")
    assert r.sucesso, r.erro
    assert (r.recebidos, r.novos, r.fora_do_municipio, r.invalidos) == (3, 2, 1, 1)
    alerta = AlertaAmbiental.objects.get(id_externo="F1")
    assert alerta.status == "detectado" and alerta.tipo == "foco_calor"
    assert alerta.bairros == [{"id": str(territorios["a"].id), "nome": "Bairro A"}]
    assert alerta.data_deteccao.isoformat().startswith("2026-10-30T17:20")
    assert (
        alerta.ingerido_em.date() != alerta.data_deteccao.date()
        or alerta.ingerido_em > alerta.data_deteccao
    )
    assert alerta.propriedades["satelite"] == "AQUA_M-T"
    # CA-ALR-006: reingerir não duplica.
    r = fontes.ingerir("inpe_focos", FOCOS.encode(), "focos.csv")
    assert (r.novos, r.ja_existentes) == (0, 2)


def test_ingestao_de_poligono_reprojetado(territorios):
    r = fontes.ingerir("inpe_deter", poligono_utm(), "deter.geojson")
    assert r.sucesso and r.novos == 1, r.erro
    alerta = AlertaAmbiental.objects.get()
    assert alerta.tipo == "supressao" and alerta.geom.geom_type == "Polygon"
    assert alerta.src_original.endswith("31984")
    assert alerta.area_ha == pytest.approx(1.0, rel=0.02)


def test_ingestao_sem_municipio_registra_erro(db):
    r = fontes.ingerir("inpe_focos", FOCOS.encode(), "focos.csv")
    assert not r.sucesso and "município" in r.erro
    assert IngestaoAlertas.objects.get().concluida_em is not None


@responses.activate
def test_tarefa_periodica_baixa_fontes_configuradas(territorios, settings):
    settings.ALERTAS_FONTES = {
        "inpe_focos": "https://fonte.teste/focos.csv",
        "inpe_deter": "https://fonte.teste/off",
    }
    responses.get("https://fonte.teste/focos.csv", body=FOCOS)
    responses.get("https://fonte.teste/off", status=503)
    resultado = ingerir_fontes_configuradas()
    assert resultado[0] == {"fonte": "inpe_focos", "sucesso": True, "novos": 2}
    assert resultado[1]["sucesso"] is False
    assert "503" in IngestaoAlertas.objects.get(fonte="inpe_deter").erro


# --- Triagem ------------------------------------------------------------------


@pytest.fixture
def alerta_a(territorios):
    fontes.ingerir("inpe_focos", FOCOS.encode(), "focos.csv")
    return AlertaAmbiental.objects.get(id_externo="F1")


def triar(cliente, alerta, **dados):
    return cliente.post(f"/api/v1/alertas/{alerta.id}/triagem", dados, format="json")


def test_fluxo_de_triagem_ate_confirmacao(cliente, analista, alerta_a):
    cliente.como(analista)
    assert triar(cliente, alerta_a, status="confirmado").status_code == 400  # salto proibido
    assert triar(cliente, alerta_a, status="em_triagem").json()["status"] == "em_triagem"
    assert (
        triar(cliente, alerta_a, status="encaminhado", justificativa="Área de APP").status_code
        == 400
    )
    r = triar(
        cliente,
        alerta_a,
        status="encaminhado",
        justificativa="Área de APP",
        referencia="SEUMA 123/2026",
    )
    assert r.json()["status"] == "encaminhado"
    assert (
        triar(cliente, alerta_a, status="confirmado", justificativa="Auto lavrado").status_code
        == 400
    )
    r = triar(
        cliente,
        alerta_a,
        status="confirmado",
        justificativa="Auto lavrado",
        referencia="Auto 45/2026",
    )
    assert r.json()["status"] == "confirmado" and r.json()["transicoes_possiveis"] == []
    historico = cliente.get(f"/api/v1/alertas/{alerta_a.id}/historico").json()
    assert historico[0]["contexto"]["motivo"] == "Auto lavrado"


def test_descartar_exige_justificativa(cliente, analista, alerta_a):
    cliente.como(analista)
    triar(cliente, alerta_a, status="em_triagem")
    assert triar(cliente, alerta_a, status="descartado").status_code == 400
    assert (
        triar(
            cliente, alerta_a, status="descartado", justificativa="Queima controlada autorizada"
        ).status_code
        == 200
    )
    assert cliente.get("/api/v1/alertas").json()["count"] == 1  # só F2 continua aberto


def test_nenhum_status_afirma_ilegalidade():
    from itree.alertas.models import StatusAlerta

    for _, rotulo in StatusAlerta.choices:
        assert "ilegal" not in rotulo.lower() and "infração" not in rotulo.lower()


def test_triagem_exige_permissao_e_escopo(cliente, gestor_municipal, tecnico_a, alerta_a):
    assert triar(cliente.como(gestor_municipal), alerta_a, status="em_triagem").status_code == 403
    # Técnico não lê alertas.
    assert cliente.como(tecnico_a).get("/api/v1/alertas").status_code == 403


def test_alerta_fora_do_escopo_nao_aparece(cliente, territorios, alerta_a):
    so_b = com_escopo(criar_usuario("b@itree.teste", ["analista_ambiental"]), territorios["b"])
    codigos = [a["id_externo"] for a in cliente.como(so_b).get("/api/v1/alertas").json()["results"]]
    assert codigos == ["F2"]


def test_conversao_em_ocorrencia(cliente, analista, gestor_municipal, alerta_a, tipo_queda):
    # CA-OCO-003: alerta não convertido não aparece nas ocorrências.
    assert cliente.como(gestor_municipal).get("/api/v1/ocorrencias").json()["count"] == 0
    analista.papeis = ["analista_ambiental", "tecnico"]
    analista.save()
    r = cliente.como(analista).post(
        f"/api/v1/alertas/{alerta_a.id}/converter-ocorrencia",
        {"tipo_id": str(tipo_queda.id), "descricao": "Verificar foco"},
        format="json",
    )
    assert r.status_code == 200, r.json()
    ocorrencia = Ocorrencia.objects.get()
    assert ocorrencia.origem == "alerta" and ocorrencia.origem_referencia == alerta_a.codigo
    assert ocorrencia.prioridade == "" and not ocorrencia.risco_confirmado
    assert (
        cliente.post(
            f"/api/v1/alertas/{alerta_a.id}/converter-ocorrencia",
            {"tipo_id": str(tipo_queda.id), "descricao": "x"},
            format="json",
        ).status_code
        == 400
    )


def test_upload_manual_e_estado_das_ingestoes(cliente, admin, analista, territorios):
    arquivo = SimpleUploadedFile("focos.csv", FOCOS.encode())
    assert (
        cliente.como(analista)
        .post(
            "/api/v1/alertas/ingestoes/arquivo",
            {"fonte": "inpe_focos", "arquivo": arquivo},
            format="multipart",
        )
        .status_code
        == 403
    )
    arquivo.seek(0)
    r = cliente.como(admin).post(
        "/api/v1/alertas/ingestoes/arquivo",
        {"fonte": "inpe_focos", "arquivo": arquivo},
        format="multipart",
    )
    assert r.status_code == 201 and r.json()["novos"] == 2
    assert cliente.como(analista).get("/api/v1/alertas/ingestoes").json()[0]["usuario"] == "admin"


def test_mapa_de_alertas(cliente, analista, alerta_a):
    corpo = cliente.como(analista).get("/api/v1/mapa/alertas").json()
    assert {f["properties"]["codigo"] for f in corpo["features"]} == {
        a.codigo for a in AlertaAmbiental.objects.all()
    }


# --- Relatórios ---------------------------------------------------------------


@pytest.fixture
def inventario(cliente, revisor):
    for ponto in (PONTO_A, PONTO_A, PONTO_B):
        cliente.como(revisor).post("/api/v1/arvores", {"localizacao": ponto}, format="json")


@pytest.fixture
def gestor_exporta(gestor_municipal):
    return gestor_municipal  # gestor tem relatorios.exportar


def test_exportacao_csv_com_metadados_e_total_igual_a_lista(
    cliente, gestor_exporta, inventario, territorios
):
    filtros = {"bairro": str(territorios["a"].id)}
    total = cliente.como(gestor_exporta).get("/api/v1/arvores", filtros).json()["count"]
    r = cliente.get(
        "/api/v1/relatorios/exportar", {"conjunto": "arvores", "formato": "csv", **filtros}
    )
    assert r.status_code == 200
    texto = r.content.decode("utf-8-sig")
    metadados = [linha for linha in texto.splitlines() if linha.startswith("#")]
    assert any("gerado_por" in m for m in metadados) and any("filtros" in m for m in metadados)
    dados = [linha for linha in texto.splitlines() if not linha.startswith("#")]
    assert len(dados) - 1 == total == 2  # CA-REL-001
    assert Exportacao.objects.get().linhas == 2


def test_exportacao_xlsx_e_geojson(cliente, gestor_exporta, inventario):
    cliente.como(gestor_exporta)
    r = cliente.get("/api/v1/relatorios/exportar", {"conjunto": "arvores", "formato": "xlsx"})
    livro = load_workbook(io.BytesIO(r.content))
    assert livro.sheetnames == ["Dados", "Metadados"]
    assert livro["Dados"].max_row == 4
    r = cliente.get("/api/v1/relatorios/exportar", {"conjunto": "arvores", "formato": "geojson"})
    corpo = json.loads(r.content)
    assert corpo["metadata"]["sistema_de_referencia"].startswith("EPSG:4326")
    assert len(corpo["features"]) == 3


def test_exportacao_exige_permissao(cliente, tecnico_a):
    r = cliente.como(tecnico_a).get(
        "/api/v1/relatorios/exportar", {"conjunto": "arvores", "formato": "csv"}
    )
    assert r.status_code == 403


def test_analise_por_bairro_e_pdf(cliente, gestor_exporta, inventario):
    corpo = (
        cliente.como(gestor_exporta)
        .get("/api/v1/relatorios/analise", {"analise": "arvores_por_bairro"})
        .json()
    )
    assert {linha["rotulo"]: linha["valor"] for linha in corpo["linhas"]} == {
        "Bairro A": 2,
        "Bairro B": 1,
    }
    assert corpo["por_bairro"] and corpo["definicao"]
    r = cliente.get("/api/v1/relatorios/painel.pdf")
    assert r.status_code == 200 and r.content.startswith(b"%PDF")
    assert cliente.get("/api/v1/relatorios/analise", {"analise": "x"}).status_code == 404


# --- Qualidade e IA -----------------------------------------------------------


def test_pendencias_de_qualidade(cliente, revisor, inventario):
    arvore = Arvore.objects.first()
    Medida.objects.create(
        arvore=arvore, tipo="altura", valor=95, unidade="m", metodo="trena", origem="medido"
    )
    Arvore.objects.create(origem_registro="importacao", especie_texto_original="Oity")
    pend = {
        p["id"]: p
        for p in cliente.como(revisor).get("/api/v1/qualidade/pendencias").json()["pendencias"]
    }
    assert pend["medidas_implausiveis"]["quantidade"] == 1
    assert pend["medidas_implausiveis"]["exemplos"][0]["valor"] == 95.0
    assert pend["sem_localizacao"]["quantidade"] == 1
    assert pend["especie_fora_catalogo"]["quantidade"] == 1
    assert all(p["criterio"] for p in pend.values())


def test_sugestao_de_ia_nao_altera_cadastro_sem_decisao(
    cliente, admin, revisor, tecnico_a, inventario
):
    arvore = Arvore.objects.first()
    oiti = Especie.objects.create(nome_cientifico="Licania tomentosa")
    corpo = {
        "tarefa": "especie",
        "arvore": str(arvore.id),
        "modelo": "classificador-teste",
        "versao_modelo": "0.1",
        "resultado": {"rotulos": ["Licania tomentosa"]},
        "especies_sugeridas": [str(oiti.id)],
    }
    com_confianca = {**corpo, "resultado": {"rotulos": ["x"], "confianca": 0.93}}
    assert (
        cliente.como(admin).post("/api/v1/inferencias", com_confianca, format="json").status_code
        == 400
    )
    r = cliente.post("/api/v1/inferencias", corpo, format="json")
    assert r.status_code == 201, r.json()
    inferencia = r.json()
    arvore.refresh_from_db()
    assert arvore.especie_id is None  # RN-INV-003

    url = f"/api/v1/inferencias/{inferencia['id']}/decidir"
    assert (
        cliente.como(tecnico_a).post(url, {"decisao": "aceita"}, format="json").status_code == 403
    )
    assert cliente.como(revisor).post(url, {"decisao": "aceita"}, format="json").status_code == 400
    r = cliente.post(url, {"decisao": "aceita", "especie": str(oiti.id)}, format="json")
    assert r.json()["decisao"] == "aceita" and r.json()["decidida_por_nome"] == "revisor"
    arvore.refresh_from_db()
    assert (arvore.especie_id, arvore.especie_estado) == (oiti.id, "confirmada")
    assert InferenciaIA.objects.get().resultado == {
        "rotulos": ["Licania tomentosa"]
    }  # original preservado
