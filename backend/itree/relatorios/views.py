"""Análises e exportações (W-09). Todo arquivo carrega metadados: data de
geração, usuário, filtros, definições e fontes (RF-REL-003)."""

import csv
import io
import json
from datetime import date, timedelta

from django.conf import settings
from django.db.models import Avg, Count, F
from django.db.models.functions import TruncMonth
from django.http import HttpResponse
from django.utils import timezone
from drf_spectacular.types import OpenApiTypes
from drf_spectacular.utils import OpenApiParameter, extend_schema, inline_serializer
from rest_framework import serializers
from rest_framework.decorators import api_view
from rest_framework.exceptions import NotFound, PermissionDenied, ValidationError
from rest_framework.request import Request
from rest_framework.response import Response

from itree.alertas.models import AlertaAmbiental
from itree.contas.acesso import filtrar_por_escopo
from itree.inspecoes.models import EstadoInspecao, Inspecao
from itree.inventario.models import Arvore, SituacaoArvore
from itree.inventario.views import filtrar_arvores
from itree.ocorrencias.models import EstadoOcorrencia, Ocorrencia
from itree.ocorrencias.views import filtrar_ocorrencias
from itree.territorio.models import Territorio, TipoTerritorio

from .models import Exportacao

FONTE = "Engine de Inteligência Arbórea — registros da plataforma"


def _periodo(p) -> tuple[date, date]:
    hoje = timezone.localdate()
    try:
        fim = date.fromisoformat(p["fim"]) if p.get("fim") else hoje
        inicio = date.fromisoformat(p["inicio"]) if p.get("inicio") else fim - timedelta(days=365)
    except ValueError as exc:
        raise ValidationError({"periodo": ["Use AAAA-MM-DD."]}) from exc
    return inicio, fim


# --- Análises -----------------------------------------------------------------

ANALISES = {
    "arvores_por_bairro": {
        "titulo": "Árvores com registro válido por bairro",
        "definicao": "Árvores ativas, com localização e origem conhecida, agrupadas pelo bairro que contém o ponto.",
        "permissao": "inventario.ler",
        "por_bairro": True,
    },
    "ocorrencias_por_tipo": {
        "titulo": "Ocorrências registradas por tipo no período",
        "definicao": "Contagem por tipo, pela data de registro. Não inclui alertas ambientais não convertidos.",
        "permissao": "ocorrencias.ler",
    },
    "ocorrencias_por_estado": {
        "titulo": "Ocorrências por situação atual",
        "definicao": "Situação atual das ocorrências registradas no período.",
        "permissao": "ocorrencias.ler",
    },
    "tempo_atendimento_por_tipo": {
        "titulo": "Tempo médio de atendimento por tipo (dias)",
        "definicao": "Média entre registro e conclusão das ocorrências concluídas no período. Tipos sem conclusões não aparecem.",
        "permissao": "ocorrencias.ler",
    },
    "inspecoes_por_mes": {
        "titulo": "Inspeções realizadas por mês",
        "definicao": "Inspeções enviadas para revisão ou validadas, pela data da inspeção.",
        "permissao": "inspecoes.ler",
    },
    "alertas_por_fonte": {
        "titulo": "Alertas ambientais por fonte e status",
        "definicao": "Alertas detectados no período. Alertas são sinais para análise, não confirmação de infração.",
        "permissao": "alertas.ler",
    },
}


def _linhas_analise(chave: str, usuario, inicio: date, fim: date) -> list[dict]:
    ocorrencias = filtrar_por_escopo(Ocorrencia.objects.all(), usuario).filter(
        criado_em__date__gte=inicio, criado_em__date__lte=fim
    )
    if chave == "arvores_por_bairro":
        validas = filtrar_por_escopo(
            Arvore.objects.filter(situacao=SituacaoArvore.ATIVA, geom__isnull=False).exclude(
                localizacao_origem=""
            ),
            usuario,
        )
        contagem = dict(validas.values_list("bairro_id").annotate(n=Count("id")))
        linhas = [
            {"rotulo": b.nome, "valor": contagem.get(b.id, 0), "territorio_id": str(b.id)}
            for b in Territorio.objects.filter(tipo=TipoTerritorio.BAIRRO).order_by("nome")
        ]
        if contagem.get(None):
            linhas.append(
                {
                    "rotulo": "Fora dos bairros cadastrados",
                    "valor": contagem[None],
                    "territorio_id": None,
                }
            )
        return sorted(linhas, key=lambda linha: -linha["valor"])
    if chave == "ocorrencias_por_tipo":
        return [
            {"rotulo": r["tipo__nome"], "valor": r["n"]}
            for r in ocorrencias.values("tipo__nome").annotate(n=Count("id")).order_by("-n")
        ]
    if chave == "ocorrencias_por_estado":
        rotulos = dict(EstadoOcorrencia.choices)
        return [
            {"rotulo": rotulos[r["estado"]], "valor": r["n"]}
            for r in ocorrencias.values("estado").annotate(n=Count("id")).order_by("-n")
        ]
    if chave == "tempo_atendimento_por_tipo":
        concluidas = filtrar_por_escopo(Ocorrencia.objects.all(), usuario).filter(
            estado=EstadoOcorrencia.CONCLUIDA,
            concluida_em__date__gte=inicio,
            concluida_em__date__lte=fim,
        )
        return [
            {
                "rotulo": r["tipo__nome"],
                "valor": round(r["media"].total_seconds() / 86400, 1),
                "base": r["n"],
            }
            for r in concluidas.values("tipo__nome")
            .annotate(media=Avg(F("concluida_em") - F("criado_em")), n=Count("id"))
            .order_by("tipo__nome")
        ]
    if chave == "inspecoes_por_mes":
        inspecoes = filtrar_por_escopo(
            Inspecao.objects.all(), usuario, campo_geom="arvore__geom"
        ).filter(
            estado__in=[EstadoInspecao.EM_REVISAO, EstadoInspecao.VALIDADA],
            realizada_em__date__gte=inicio,
            realizada_em__date__lte=fim,
        )
        return [
            {"rotulo": r["mes"].strftime("%m/%Y"), "valor": r["n"]}
            for r in inspecoes.annotate(mes=TruncMonth("realizada_em"))
            .values("mes")
            .annotate(n=Count("id"))
            .order_by("mes")
        ]
    if chave == "alertas_por_fonte":
        alertas = filtrar_por_escopo(AlertaAmbiental.objects.all(), usuario).filter(
            data_deteccao__date__gte=inicio, data_deteccao__date__lte=fim
        )
        fontes = dict(AlertaAmbiental._meta.get_field("fonte").choices)
        status = dict(AlertaAmbiental._meta.get_field("status").choices)
        return [
            {"rotulo": f"{fontes[r['fonte']]} — {status[r['status']]}", "valor": r["n"]}
            for r in alertas.values("fonte", "status")
            .annotate(n=Count("id"))
            .order_by("fonte", "status")
        ]
    raise NotFound("Análise desconhecida.")


@extend_schema(
    operation_id="relatorios_analise",
    parameters=[
        OpenApiParameter("analise", str, required=True, description=", ".join(ANALISES)),
        OpenApiParameter("inicio", str),
        OpenApiParameter("fim", str),
    ],
    responses=inline_serializer(
        "Analise",
        {
            "id": serializers.CharField(),
            "titulo": serializers.CharField(),
            "definicao": serializers.CharField(),
            "fonte": serializers.CharField(),
            "periodo": serializers.DictField(),
            "calculado_em": serializers.DateTimeField(),
            "por_bairro": serializers.BooleanField(),
            "linhas": serializers.ListField(child=serializers.DictField()),
        },
    ),
)
@api_view(["GET"])
def analise(request: Request):
    chave = request.query_params.get("analise", "")
    definicao = ANALISES.get(chave)
    if definicao is None:
        raise NotFound("Análise desconhecida.")
    if not (
        request.user.tem_permissao("relatorios.ler")
        and request.user.tem_permissao(definicao["permissao"])
    ):
        raise PermissionDenied()
    inicio, fim = _periodo(request.query_params)
    return Response(
        {
            "id": chave,
            "titulo": definicao["titulo"],
            "definicao": definicao["definicao"],
            "fonte": FONTE,
            "periodo": {"inicio": inicio, "fim": fim},
            "calculado_em": timezone.now(),
            "por_bairro": definicao.get("por_bairro", False),
            "linhas": _linhas_analise(chave, request.user, inicio, fim),
        }
    )


@extend_schema(
    operation_id="relatorios_catalogo",
    responses=inline_serializer(
        "CatalogoAnalises", {"analises": serializers.ListField(child=serializers.DictField())}
    ),
)
@api_view(["GET"])
def catalogo(request: Request):
    return Response(
        {
            "analises": [
                {"id": k, "titulo": v["titulo"], "definicao": v["definicao"]}
                for k, v in ANALISES.items()
                if request.user.tem_permissao(v["permissao"])
            ]
        }
    )


# --- Exportações --------------------------------------------------------------


def _dados_arvores(usuario, params):
    qs = filtrar_arvores(
        filtrar_por_escopo(
            Arvore.objects.select_related("especie", "bairro"), usuario, incluir_sem_geom=True
        ),
        params,
    )
    colunas = [
        "codigo",
        "especie",
        "especie_estado",
        "latitude",
        "longitude",
        "localizacao_origem",
        "precisao_m",
        "endereco",
        "bairro",
        "condicao_aparente",
        "ultima_inspecao_em",
        "situacao",
        "origem_registro",
    ]

    def linha(a):
        return {
            "codigo": a.codigo,
            "especie": a.especie.nome_cientifico if a.especie else a.especie_texto_original,
            "especie_estado": a.especie_estado,
            "latitude": a.geom.y if a.geom else None,
            "longitude": a.geom.x if a.geom else None,
            "localizacao_origem": a.localizacao_origem,
            "precisao_m": a.precisao_m,
            "endereco": a.endereco,
            "bairro": a.bairro.nome if a.bairro else "",
            "condicao_aparente": a.condicao_aparente,
            "ultima_inspecao_em": a.ultima_inspecao_em.isoformat() if a.ultima_inspecao_em else "",
            "situacao": a.situacao,
            "origem_registro": a.origem_registro,
        }

    return colunas, [(linha(a), a.geom) for a in qs.order_by("numero")[: LIMITE_EXPORTACAO + 1]]


def _dados_ocorrencias(usuario, params):
    qs = filtrar_ocorrencias(
        filtrar_por_escopo(Ocorrencia.objects.select_related("tipo", "bairro", "arvore"), usuario),
        params,
    )
    colunas = [
        "codigo",
        "tipo",
        "estado",
        "prioridade",
        "risco_confirmado",
        "condicao_aparente",
        "descricao",
        "arvore",
        "bairro",
        "latitude",
        "longitude",
        "prazo",
        "registrada_em",
        "concluida_em",
        "origem",
    ]

    def linha(o):
        return {
            "codigo": o.codigo,
            "tipo": o.tipo.nome,
            "estado": o.estado,
            "prioridade": o.prioridade,
            "risco_confirmado": o.risco_confirmado,
            "condicao_aparente": o.condicao_aparente,
            "descricao": o.descricao,
            "arvore": o.arvore.codigo if o.arvore else "",
            "bairro": o.bairro.nome if o.bairro else "",
            "latitude": o.geom.y,
            "longitude": o.geom.x,
            "prazo": o.prazo.isoformat() if o.prazo else "",
            "registrada_em": o.criado_em.isoformat(),
            "concluida_em": o.concluida_em.isoformat() if o.concluida_em else "",
            "origem": o.origem,
        }

    return colunas, [(linha(o), o.geom) for o in qs.order_by("numero")[: LIMITE_EXPORTACAO + 1]]


def _dados_alertas(usuario, params):
    qs = filtrar_por_escopo(AlertaAmbiental.objects.all(), usuario)
    if s := params.get("status"):
        qs = qs.filter(status__in=s.split(","))
    colunas = [
        "codigo",
        "fonte",
        "id_externo",
        "tipo",
        "data_deteccao",
        "ingerido_em",
        "area_ha",
        "status",
        "bairros",
    ]

    def linha(a):
        return {
            "codigo": a.codigo,
            "fonte": a.fonte,
            "id_externo": a.id_externo,
            "tipo": a.tipo,
            "data_deteccao": a.data_deteccao.isoformat(),
            "ingerido_em": a.ingerido_em.isoformat(),
            "area_ha": a.area_ha,
            "status": a.status,
            "bairros": ", ".join(b["nome"] for b in a.bairros),
        }

    return colunas, [(linha(a), a.geom) for a in qs.order_by("numero")[: LIMITE_EXPORTACAO + 1]]


CONJUNTOS = {
    "arvores": ("Inventário de árvores", "inventario.ler", _dados_arvores),
    "ocorrencias": ("Ocorrências", "ocorrencias.ler", _dados_ocorrencias),
    "alertas": (
        "Alertas ambientais (sinais para análise, não confirmação de infração)",
        "alertas.ler",
        _dados_alertas,
    ),
}
FORMATOS = {
    "csv": "text/csv; charset=utf-8",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "geojson": "application/geo+json",
}


# Linhas por arquivo. Busca-se uma a mais para saber se houve corte.
LIMITE_EXPORTACAO = 50_000

# Texto que planilhas interpretariam como fórmula (CWE-1236).
PREFIXOS_FORMULA = ("=", "+", "-", "@", "\t", "\r")


def _texto_seguro_csv(valor):
    """No CSV, prefixa com apóstrofo o texto que viraria fórmula ao abrir no Excel."""
    if isinstance(valor, str) and valor.startswith(PREFIXOS_FORMULA):
        return "'" + valor
    return valor


def _celula(planilha, valor):
    """No XLSX, grava texto sempre como texto (o openpyxl trata '=...' como fórmula)."""
    from openpyxl.cell import WriteOnlyCell

    celula = WriteOnlyCell(planilha, value=valor)
    if isinstance(valor, str):
        celula.data_type = "s"
    return celula


def _metadados(usuario, titulo: str, params: dict, linhas: int) -> dict:
    return {
        "titulo": titulo,
        "gerado_em": timezone.now().isoformat(),
        "gerado_por": usuario.nome or usuario.email,
        "filtros": params,
        "registros": linhas,
        "fonte": FONTE,
        "sistema_de_referencia": "EPSG:4326 (WGS 84)",
        "ambiente": settings.AMBIENTE,
        "aviso": "DADOS DE DEMONSTRAÇÃO — não representam a realidade."
        if settings.AMBIENTE == "demonstracao"
        else "",
        "observacao": "Valores marcados como estimados ou importados não equivalem a medições de campo.",
    }


@extend_schema(
    operation_id="relatorios_exportar",
    parameters=[
        OpenApiParameter(
            "conjunto", str, required=True, description="arvores | ocorrencias | alertas"
        ),
        OpenApiParameter("formato", str, required=True, description="csv | xlsx | geojson"),
    ],
    responses={(200, "application/octet-stream"): OpenApiTypes.BINARY},
)
@api_view(["GET"])
def exportar(request: Request):
    p = {k: v for k, v in request.query_params.items() if k not in ("conjunto", "formato")}
    conjunto, formato = request.query_params.get("conjunto"), request.query_params.get("formato")
    if conjunto not in CONJUNTOS:
        raise ValidationError({"conjunto": ["Conjunto desconhecido."]})
    if formato not in FORMATOS:
        raise ValidationError({"formato": ["Use csv, xlsx ou geojson."]})
    titulo, permissao, fonte_dados = CONJUNTOS[conjunto]
    if not (
        request.user.tem_permissao("relatorios.exportar") and request.user.tem_permissao(permissao)
    ):
        raise PermissionDenied()
    colunas, registros = fonte_dados(request.user, p)
    truncado = len(registros) > LIMITE_EXPORTACAO
    registros = registros[:LIMITE_EXPORTACAO]
    meta = _metadados(request.user, titulo, p, len(registros))
    if truncado:
        meta["truncado"] = (
            f"O resultado passou de {LIMITE_EXPORTACAO} registros e foi cortado; "
            "refine os filtros para exportar o restante."
        )
    Exportacao.objects.create(
        usuario=request.user, conjunto=conjunto, formato=formato, filtros=p, linhas=len(registros)
    )
    nome = f"itree-{conjunto}-{timezone.localdate():%Y%m%d}.{formato}"

    if formato == "csv":
        saida = io.StringIO()
        for chave, valor in meta.items():
            if valor not in ("", None, {}):
                texto = json.dumps(valor, ensure_ascii=False) if isinstance(valor, dict) else valor
                saida.write(f"# {chave}: {str(texto).replace(chr(10), ' ')}\n")
        escritor = csv.DictWriter(saida, fieldnames=colunas, delimiter=";")
        escritor.writeheader()
        for linha, _ in registros:
            escritor.writerow({c: _texto_seguro_csv(v) for c, v in linha.items()})
        conteudo = ("﻿" + saida.getvalue()).encode("utf-8")
    elif formato == "xlsx":
        from openpyxl import Workbook

        # write_only grava em fluxo, sem um objeto por célula na memória.
        livro = Workbook(write_only=True)
        planilha = livro.create_sheet("Dados")
        planilha.append(colunas)
        for linha, _ in registros:
            planilha.append([_celula(planilha, linha[c]) for c in colunas])
        info = livro.create_sheet("Metadados")
        for chave, valor in meta.items():
            texto = json.dumps(valor, ensure_ascii=False) if isinstance(valor, dict) else valor
            info.append([chave, _celula(info, texto)])
        buffer = io.BytesIO()
        livro.save(buffer)
        conteudo = buffer.getvalue()
    else:
        conteudo = json.dumps(
            {
                "type": "FeatureCollection",
                "metadata": meta,
                "features": [
                    {
                        "type": "Feature",
                        "geometry": json.loads(geom.geojson) if geom else None,
                        "properties": linha,
                    }
                    for linha, geom in registros
                ],
            },
            ensure_ascii=False,
            default=str,
        ).encode("utf-8")

    resposta = HttpResponse(conteudo, content_type=FORMATOS[formato])
    resposta["Content-Disposition"] = f'attachment; filename="{nome}"'
    return resposta


@extend_schema(
    operation_id="relatorios_painel_pdf",
    parameters=[
        OpenApiParameter("inicio", str),
        OpenApiParameter("fim", str),
        OpenApiParameter("bairro", str),
    ],
    responses={(200, "application/pdf"): OpenApiTypes.BINARY},
)
@api_view(["GET"])
def painel_pdf(request: Request):
    """Relatório em PDF dos indicadores, com definição, período e fonte (CA-REL-002)."""
    from fpdf import FPDF

    from itree.indicadores.views import painel as calcular_painel

    if not request.user.tem_permissao("relatorios.exportar"):
        raise PermissionDenied()
    dados = calcular_painel(request._request).data
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 15)
    pdf.cell(0, 9, "Engine de Inteligencia Arborea - Indicadores", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", size=9)

    def bloco(altura: float, conteudo: str) -> None:
        pdf.multi_cell(0, altura, conteudo, new_x="LMARGIN", new_y="NEXT")

    def texto(t: str) -> str:
        # Fonte padrão do PDF é Latin-1.
        return str(t).encode("latin-1", "replace").decode("latin-1")

    bloco(
        5,
        texto(
            f"Periodo: {dados['periodo']['inicio']} a {dados['periodo']['fim']} | "
            f"Filtros: bairro={dados['filtros']['bairro'] or 'todos'}, equipe={dados['filtros']['equipe'] or 'todas'}\n"
            f"Gerado em {timezone.localtime():%d/%m/%Y %H:%M} por {request.user.nome or request.user.email}. Fonte: {FONTE}."
        ),
    )
    if settings.AMBIENTE == "demonstracao":
        pdf.set_text_color(160, 90, 0)
        bloco(6, "DADOS DE DEMONSTRACAO - nao representam a realidade.")
        pdf.set_text_color(0, 0, 0)
    pdf.ln(3)
    for i in dados["indicadores"]:
        valor = (
            "Sem dados"
            if i["valor"] is None
            else f"{i['valor']}{' ' + i['unidade'] if i['unidade'] else ''}"
        )
        pdf.set_font("Helvetica", "B", 11)
        bloco(6, texto(f"{i['nome']}: {valor}"))
        pdf.set_font("Helvetica", size=9)
        bloco(
            5,
            texto(
                f"{i['definicao']} Formula: {i['formula']}. Periodo: {i['periodo']}. Fonte: {i['fonte']}."
            ),
        )
        pdf.ln(1.5)
    Exportacao.objects.create(
        usuario=request.user,
        conjunto="painel",
        formato="pdf",
        filtros=dict(request.query_params.items()),
        linhas=len(dados["indicadores"]),
    )
    resposta = HttpResponse(bytes(pdf.output()), content_type="application/pdf")
    resposta["Content-Disposition"] = (
        f'attachment; filename="itree-indicadores-{timezone.localdate():%Y%m%d}.pdf"'
    )
    return resposta
